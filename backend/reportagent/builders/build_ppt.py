# -*- coding: utf-8 -*-
"""slides.json → PPTX 조립.

사용법: python scripts/build_ppt.py <slides.json> [output_dir]
기본 출력 디렉토리는 입력 slides.json의 2단 위 디렉토리/output
(예: projects/X/work/slides.json → projects/X/output).
스키마는 .claude/skills/ppt-design/SKILL.md 참조.

템플릿: meta.template > <프로젝트>/assets/template.pptx > templates/template.pptx
순으로 찾아 빈 레이아웃에 배치한다 (16:9가 아니면 경고 후 무시).
meta.template에 "none"을 주면 템플릿을 강제로 쓰지 않는다.
"""
import json
import re
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt
except ImportError:
    print("python-pptx가 필요합니다: pip install python-pptx")
    sys.exit(1)

sys.path.insert(0, str(Path(__file__).parent))
import theme as T

REQUIRED = {
    "cover": {"title"},
    "toc": {"title", "bullets"},
    "two-col": {"title", "left", "right"},
    "table": {"title", "table"},
    "chart": {"title", "chart"},
    "arch": {"title", "arch"},
    "closing": {"title", "bullets"},
}
TYPES = set(REQUIRED)

SLIDE_NO_TOP = 7.0
CONTENT_W = 12.1  # 슬라이드 내용 폭 (SLIDE_W - 2*MARGIN)
TEMPLATE_FILENAME = "template.pptx"
BLANK_LAYOUT_IDX = 6  # build()에서 템플릿의 빈 레이아웃 인덱스로 갱신됨


def _hex(color: str):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(color)


def _set_text(tf, runs, size, color=T.TEXT, bold=False, align=PP_ALIGN.LEFT):
    """runs: str 또는 [(text, bold), ...] 리스트. 첫 단락(paragraph) 재사용."""
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    if isinstance(runs, str):
        runs = [(runs, False)]
    for i, (txt, b) in enumerate(runs):
        r = p.add_run()
        r.text = txt
        r.font.name = T.FONT
        r.font.size = Pt(size)
        r.font.bold = bold or b
        r.font.color.rgb = _hex(color)


def _add_text(slide, pos, size_wh, runs, size, color=T.TEXT, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(pos[0]), Inches(pos[1]), Inches(size_wh[0]), Inches(size_wh[1]))
    _set_text(box.text_frame, runs, size, color, bold, align)
    return box


def _add_rect(slide, pos, size_wh, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(pos[0]), Inches(pos[1]), Inches(size_wh[0]), Inches(size_wh[1]))
    sh.fill.solid()
    sh.fill.fore_color.rgb = _hex(fill)
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def _bullets_into(tf, bullets, size=13, color=T.TEXT):
    """bullets: [{label, body}] → label 굵게 + body."""
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        label = (b.get("label") or "").strip()
        body = b.get("body", "")
        if label:
            r = p.add_run()
            r.text = f"{label}  "
            r.font.name = T.FONT
            r.font.size = Pt(size)
            r.font.bold = True
            r.font.color.rgb = _hex(T.BLUE if color == T.TEXT else color)
        r2 = p.add_run()
        r2.text = body
        r2.font.name = T.FONT
        r2.font.size = Pt(size)
        r2.font.color.rgb = _hex(color)
        p.space_after = Pt(8)


def _add_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT_IDX])


def _resolve_template(jpath: Path, meta: dict):
    """템플릿 경로 결정: meta.template(경로) > <프로젝트>/assets/template.pptx > templates/template.pptx > 없음.
    meta.template이 "none"이면 템플릿을 강제로 쓰지 않는다 (스모크 테스트·템플릿 무시용)."""
    override = str(meta.get("template", "") or "").strip()
    if override.lower() == "none":
        return None
    if override:
        return Path(override)
    project_dir = jpath.resolve().parent.parent
    for cand in (project_dir / "assets" / TEMPLATE_FILENAME,
                 Path(__file__).resolve().parent.parent / "templates" / TEMPLATE_FILENAME):
        if cand.is_file():
            return cand
    return None


def _blank_layout_index(prs) -> int:
    """템플릿 내 placeholder가 없는(빈) 레이아웃 인덱스. 없으면 마지막 레이아웃."""
    for i, layout in enumerate(prs.slide_layouts):
        if len(layout.placeholders) == 0:
            return i
    return len(prs.slide_layouts) - 1


def _slide_title(slide, title):
    _add_text(slide, T.TITLE_POS, (CONTENT_W, 0.8), title, T.SIZE_SLIDE_TITLE, bold=True)
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(T.MARGIN), Inches(T.TITLE_DIVIDER_Y), Inches(CONTENT_W), Pt(1))
    ln.fill.solid()
    ln.fill.fore_color.rgb = _hex(T.LINE)
    ln.line.fill.background()
    ln.shadow.inherit = False


def _page_no(slide, n):
    _add_text(slide, (12.3, SLIDE_NO_TOP), (0.7, 0.3), str(n), T.SIZE_CAPTION, T.TEXT_SUB, align=PP_ALIGN.RIGHT)


# ---------- 유형별 렌더러 ----------

def render_cover(prs, slide_data, _n):
    s = _add_slide(prs)
    _add_rect(s, (0, 0), (T.SLIDE_W, T.SLIDE_H), T.NAVY)
    meta = slide_data.get("_meta", {})
    _add_text(s, (0.7, 0.5), (6, 0.4), meta.get("company", ""), T.SIZE_BODY, T.LIGHT_BLUE)
    _add_text(s, (0.7, 2.6), (11.5, 1.3), slide_data["title"], T.SIZE_COVER_TITLE, T.BG, bold=True)
    if slide_data.get("subtitle"):
        _add_text(s, (0.7, 3.9), (11, 0.6), slide_data["subtitle"], T.SIZE_HEADING, T.LIGHT_BLUE)
    footer = "  ".join(x for x in [meta.get("date", ""), meta.get("department", "")] if x)
    _add_text(s, (0.7, 6.5), (11, 0.4), footer, 12, T.LIGHT_BLUE)
    return s


def render_toc(prs, slide_data, n):
    s = _add_slide(prs)
    _slide_title(s, slide_data["title"])
    y = 2.2
    for b in slide_data["bullets"]:
        _add_text(s, (0.9, y), (1.0, 0.6), b.get("label", ""), 24, T.BLUE, bold=True)
        _add_text(s, (2.1, y + 0.08), (10.0, 0.6), b.get("body", ""), 20, T.TEXT)
        y += 1.2
    _page_no(s, n)
    return s


def render_two_col(prs, slide_data, n):
    s = _add_slide(prs)
    _slide_title(s, slide_data["title"])
    for i, key in enumerate(["left", "right"]):
        col = slide_data.get(key) or {}
        x = T.MARGIN + i * 6.2
        _add_rect(s, (x, T.BODY_TOP), (5.9, 5.2), T.BG_SOFT)
        head = col.get("heading", "")
        if head:
            _add_text(s, (x + 0.35, T.BODY_TOP + 0.25), (5.2, 0.5), head, 16, T.BLUE, bold=True)
        tb = s.shapes.add_textbox(Inches(x + 0.35), Inches(T.BODY_TOP + 0.85), Inches(5.2), Inches(4.1))
        _bullets_into(tb.text_frame, col.get("bullets", []))
    _page_no(s, n)
    return s


def render_table(prs, slide_data, n):
    s = _add_slide(prs)
    _slide_title(s, slide_data["title"])
    spec = slide_data["table"]
    headers, rows = spec["headers"], spec["rows"]
    n_rows, n_cols = len(rows) + 1, len(headers)
    height = min(0.5 * n_rows, 5.2)
    gfx = s.shapes.add_table(n_rows, n_cols, Inches(T.MARGIN), Inches(T.BODY_TOP), Inches(CONTENT_W), Inches(height))
    table = gfx.table
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = _hex(T.NAVY)
        cell.margin_top = Emu(0)
        _set_text(cell.text_frame, h, 13, T.BG, bold=True)
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _hex(T.BG_SOFT if r % 2 == 0 else T.BG)
            _set_text(cell.text_frame, str(val), 12, T.TEXT)
    if spec.get("note"):
        _add_text(s, (0.8, 6.6), (CONTENT_W, 0.3), spec["note"], T.SIZE_CAPTION, T.TEXT_SUB)
    _page_no(s, n)
    return s


def render_chart(prs, slide_data, n):
    s = _add_slide(prs)
    _slide_title(s, slide_data["title"])
    spec = slide_data["chart"]
    data = CategoryChartData()
    data.categories = spec["categories"]
    colors = [T.BLUE, T.LIGHT_BLUE, T.ACCENT]
    for i, series in enumerate(spec["series"]):
        data.add_series(series["name"], series["values"])
    gfx = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.6), data)
    chart = gfx.chart
    chart.has_legend = len(spec["series"]) > 1
    for i, plot_series in enumerate(chart.series):
        plot_series.format.fill.solid()
        plot_series.format.fill.fore_color.rgb = _hex(colors[i % len(colors)])
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.font.size = Pt(10)
    plot.data_labels.font.color.rgb = _hex(T.TEXT_SUB)
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    if slide_data.get("chart", {}).get("source"):
        _add_text(s, (0.8, 6.6), (11, 0.3), spec["source"], T.SIZE_CAPTION, T.TEXT_SUB)
    _page_no(s, n)
    return s


def render_closing(prs, slide_data, n):
    s = _add_slide(prs)
    _add_rect(s, (0, 0), (T.SLIDE_W, T.SLIDE_H), T.NAVY)
    _add_text(s, (0.7, 1.8), (11.5, 0.9), slide_data["title"], 32, T.BG, bold=True)
    tb = s.shapes.add_textbox(Inches(0.7), Inches(3.2), Inches(11.5), Inches(2.2))
    _bullets_into(tb.text_frame, slide_data.get("bullets", []), size=16, color=T.BG)
    if slide_data.get("note"):
        _add_text(s, (0.7, 5.8), (11.5, 0.5), slide_data["note"], T.SIZE_BODY, T.LIGHT_BLUE)
    return s


def render_arch(prs, slide_data, n):
    s = _add_slide(prs)
    _slide_title(s, slide_data["title"])
    spec = slide_data["arch"]
    groups = spec["groups"]
    band_gap = 0.18
    band_h = min(1.05, (5.2 - band_gap * (len(groups) - 1)) / len(groups))
    items_x, items_w = 2.75, CONTENT_W - 2.15
    name_w = 2.0
    for g, group in enumerate(groups):
        y = T.BODY_TOP + g * (band_h + band_gap)
        _add_rect(s, (T.MARGIN, y), (name_w, band_h), T.BG_SOFT)
        name_box = s.shapes.add_textbox(Inches(T.MARGIN), Inches(y), Inches(name_w), Inches(band_h))
        _set_text(name_box.text_frame, group["name"], 13, T.BLUE, bold=True, align=PP_ALIGN.CENTER)
        name_box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        items = group["items"]
        gap = 0.12
        box_w = (items_w - gap * (len(items) - 1)) / len(items)
        box_h = band_h - 0.2
        for i, item in enumerate(items):
            bx = items_x + i * (box_w + gap)
            by = y + 0.1
            rect = _add_rect(s, (bx, by), (box_w, box_h), T.BG)
            rect.line.color.rgb = _hex(T.LINE)
            rect.line.width = Pt(1)
            tb = s.shapes.add_textbox(Inches(bx), Inches(by), Inches(box_w), Inches(box_h))
            _set_text(tb.text_frame, item, 12, T.TEXT, align=PP_ALIGN.CENTER)
            tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    if spec.get("note"):
        _add_text(s, (0.8, 6.6), (CONTENT_W, 0.3), spec["note"], T.SIZE_CAPTION, T.TEXT_SUB)
    _page_no(s, n)
    return s


RENDERERS = {
    "cover": render_cover,
    "toc": render_toc,
    "two-col": render_two_col,
    "table": render_table,
    "chart": render_chart,
    "arch": render_arch,
    "closing": render_closing,
}


def _sanitize_title(title: str) -> str:
    """제목 → 파일명 접두어. Windows 금지 문자·glob 메타문자 [ ] 제거, 공백→_. (build_doc.py와 동일 규칙)"""
    for ch in '\\/:*?"<>|[]\n\r\t':
        title = title.replace(ch, " ")
    return "_".join(title.split()).strip(" ._") or "제목없음"


def _next_version(out_dir: Path, prefix: str) -> int:
    """out_dir에서 같은 접두어의 최대 vNN을 찾아 +1."""
    mx = 0
    for f in out_dir.glob(f"{prefix}_v*.pptx"):
        m = re.search(r"_v(\d+)\.pptx$", f.name, re.IGNORECASE)
        if m:
            mx = max(mx, int(m.group(1)))
    return mx + 1


def _output_filename(meta: dict, out_dir: Path) -> str:
    """output_name이 있으면 그대로(확장자는 보정 — build_doc.py와 동일 규칙), 없으면 <제목>_vNN 자동 생성."""
    if meta.get("output_name"):
        name = str(meta["output_name"])
        if name.lower().endswith(".pptx"):
            name = name[:-5]
        return f"{name}.pptx"
    prefix = _sanitize_title(meta.get("title", ""))
    return f"{prefix}_v{_next_version(out_dir, prefix):02d}.pptx"


def validate(doc):
    if not isinstance(doc, dict):
        raise ValueError("slides.json 최상위는 객체(dict)여야 합니다")
    slides = doc.get("slides", [])
    if not isinstance(slides, list) or not slides:
        raise ValueError("slides는 비어있지 않은 배열이어야 합니다")
    for i, sd in enumerate(slides, 1):
        if not isinstance(sd, dict):
            raise ValueError(f"슬라이드 {i}: 객체(dict)여야 합니다")
        t = sd.get("type")
        if t not in TYPES:
            raise ValueError(f"슬라이드 {i}: 알 수 없는 유형 '{t}' (7종 중 하나여야 함)")
        missing = REQUIRED[t] - set(sd)
        if missing:
            raise ValueError(f"슬라이드 {i} ({t}): 필수 키 누락 {missing}")
        _validate_spec(i, t, sd)
    meta = doc.get("meta", {})
    if not isinstance(meta, dict):
        raise ValueError("meta는 객체(dict)여야 합니다")
    if not (meta.get("title") or meta.get("output_name")):
        raise ValueError("meta.title 필요 (meta.output_name은 선택적 파일명 오버라이드)")
    return slides


def _validate_spec(i, t, sd):
    """중첩 스펙의 실제 렌더러 요구사항 검증 (검증 통과 후 KeyError/IndexError 방지)."""
    if t == "table":
        spec = sd["table"]
        if not (isinstance(spec, dict) and isinstance(spec.get("headers"), list) and spec["headers"]
                and isinstance(spec.get("rows"), list)):
            raise ValueError(f"슬라이드 {i} (table): table.headers(비어있지 않은 배열)·table.rows(배열) 필요")
        for r, row in enumerate(spec["rows"], 1):
            if not isinstance(row, list):
                raise ValueError(f"슬라이드 {i} (table): rows[{r}]는 배열이어야 합니다")
            if len(row) > len(spec["headers"]):
                raise ValueError(f"슬라이드 {i} (table): rows[{r}] 셀 수가 headers 수({len(spec['headers'])})를 초과합니다")
    elif t == "chart":
        spec = sd["chart"]
        if not (isinstance(spec, dict) and isinstance(spec.get("categories"), list)
                and isinstance(spec.get("series"), list) and spec["series"]):
            raise ValueError(f"슬라이드 {i} (chart): chart.categories(배열)·chart.series(비어있지 않은 배열) 필요")
        for si, sr in enumerate(spec["series"], 1):
            if not (isinstance(sr, dict) and isinstance(sr.get("name"), str)
                    and isinstance(sr.get("values"), list)):
                raise ValueError(f"슬라이드 {i} (chart): series[{si}]는 name(문자열)·values(배열)을 가져야 합니다")
    elif t == "arch":
        spec = sd["arch"]
        if not (isinstance(spec, dict) and isinstance(spec.get("groups"), list) and spec["groups"]):
            raise ValueError(f"슬라이드 {i} (arch): arch.groups(비어있지 않은 배열) 필요")
        if len(spec["groups"]) > 6:
            raise ValueError(f"슬라이드 {i} (arch): 계층은 최대 6개까지 지원합니다")
        for gi, gr in enumerate(spec["groups"], 1):
            if not (isinstance(gr, dict) and isinstance(gr.get("name"), str)
                    and isinstance(gr.get("items"), list) and gr["items"]):
                raise ValueError(f"슬라이드 {i} (arch): groups[{gi}]는 name(문자열)·items(비어있지 않은 배열)을 가져야 합니다")
            if len(gr["items"]) > 6:
                raise ValueError(f"슬라이드 {i} (arch): 계층당 구성요소는 최대 6개까지 지원합니다")


def build(json_path: str, out_dir: str = ""):
    global BLANK_LAYOUT_IDX
    jpath = Path(json_path)
    doc = json.loads(jpath.read_text(encoding="utf-8-sig"))
    slides = validate(doc)
    meta = doc.get("meta", {})

    tpath = _resolve_template(jpath, meta)
    if tpath:
        probe = Presentation(str(tpath))
        w, h = probe.slide_width / 914400, probe.slide_height / 914400
        if abs(w - T.SLIDE_W) < 0.05 and abs(h - T.SLIDE_H) < 0.05:
            prs = probe
            BLANK_LAYOUT_IDX = _blank_layout_index(prs)
            print(f"템플릿: {tpath} (빈 레이아웃 #{BLANK_LAYOUT_IDX})")
        else:
            print(f"경고: 템플릿이 16:9 ({T.SLIDE_W}×{T.SLIDE_H}in)가 아니어서 무시합니다: {tpath}")
            tpath, prs = None, Presentation()
    else:
        prs = Presentation()
    if not tpath:
        prs.slide_width = Inches(T.SLIDE_W)
        prs.slide_height = Inches(T.SLIDE_H)

    for i, sd in enumerate(slides, 1):
        sd["_meta"] = meta
        RENDERERS[sd["type"]](prs, sd, i)

    base = Path(out_dir) if out_dir else jpath.resolve().parent.parent / "output"
    base.mkdir(parents=True, exist_ok=True)
    out = base / _output_filename(meta, base)
    prs.save(str(out))
    print(f"OK: {out} ({len(slides)} slides)")


if __name__ == "__main__":
    if len(sys.argv) not in (2, 3):
        print("사용법: python scripts/build_ppt.py <slides.json> [output_dir]")
        sys.exit(1)
    try:
        build(sys.argv[1], sys.argv[2] if len(sys.argv) == 3 else "")
    except (ValueError, OSError) as e:  # JSONDecodeError·UnicodeDecodeError·FileNotFoundError 포함
        print(f"오류: {e}")
        sys.exit(1)