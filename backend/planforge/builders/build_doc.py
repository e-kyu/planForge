# -*- coding: utf-8 -*-
"""report.json → md/html/docx 조립.

사용법: python scripts/build_doc.py <report.json> <md|html|docx> [output_dir]
기본 출력 디렉토리는 입력 report.json의 2단 위 디렉토리/output
(예: projects/X/work/report.json → projects/X/output).
버전 채번은 확장자별 독립 시퀀스 ({접두어}_vNN.{ext}).
docx만 python-docx 필요 (pip install python-docx) — md/html은 표준 라이브러리만으로 동작.
스키마는 CLAUDE.md "report.json 스키마" 절 + tests/fixtures/report.sample.json 참조.
스타일은 theme.py 토큰 + 이 파일에 내장 (ppt-design 스킬은 PPT 전용).
"""
import json
import re
import sys
import html as _html
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import theme as T

FORMATS = ("md", "html", "docx")
SECTION_TYPES = {"header", "overview", "section", "conclusion"}
BLOCK_KINDS = {"prose", "table", "data"}
# 차트 계열색 — build_ppt.py와 동일 순서 (주/보조/강조)
SERIES_COLORS = (T.BLUE, T.LIGHT_BLUE, T.ACCENT)
# 문서 헤더 meta-line 색 (theme 토큰에는 없는 문서 전용 파생색)
META_LINE_COLOR = "C6D3E2"


# ---------------------------------------------------------------- 검증

def validate(doc):
    if not isinstance(doc, dict):
        raise ValueError("report.json 최상위는 객체(dict)여야 합니다")
    sections = doc.get("sections", [])
    if not isinstance(sections, list) or not sections:
        raise ValueError("sections는 비어있지 않은 배열이어야 합니다")
    meta = doc.get("meta", {})
    if not isinstance(meta, dict):
        raise ValueError("meta는 객체(dict)여야 합니다")
    if not (meta.get("title") or meta.get("output_name")):
        raise ValueError("meta.title 필요 (meta.output_name은 선택적 파일명 오버라이드)")
    for i, s in enumerate(sections, 1):
        if not isinstance(s, dict):
            raise ValueError(f"섹션 {i}: 객체(dict)여야 합니다")
        t = s.get("type")
        if t not in SECTION_TYPES:
            raise ValueError(f"섹션 {i}: 알 수 없는 유형 '{t}' ({sorted(SECTION_TYPES)} 중 하나여야 함)")
        missing = REQUIRED_SECTIONS[t] - set(s)
        if missing:
            raise ValueError(f"섹션 {i} ({t}): 필수 키 누락 {missing}")
        blocks = s.get("blocks", [])
        if not isinstance(blocks, list):
            raise ValueError(f"섹션 {i}: blocks는 배열이어야 합니다")
        for j, b in enumerate(blocks, 1):
            if not isinstance(b, dict):
                raise ValueError(f"섹션 {i} 블록 {j}: 객체(dict)여야 합니다")
            k = b.get("kind")
            if k not in BLOCK_KINDS:
                raise ValueError(f"섹션 {i} 블록 {j}: 알 수 없는 kind '{k}' ({sorted(BLOCK_KINDS)} 중 하나)")
            missing = REQUIRED_BLOCKS[k] - set(b)
            if missing:
                raise ValueError(f"섹션 {i} 블록 {j} ({k}): 필수 키 누락 {missing}")
            _validate_block(i, j, k, b)
    return sections


def _validate_block(i, j, k, b):
    """블록 내부의 실제 렌더러 요구사항 검증 (검증 통과 후 KeyError/ValueError 방지)."""
    if k == "table":
        if not (isinstance(b.get("headers"), list) and b["headers"]):
            raise ValueError(f"섹션 {i} 블록 {j} (table): headers는 비어있지 않은 배열이어야 합니다")
        for r, row in enumerate(b["rows"], 1):
            if not isinstance(row, list):
                raise ValueError(f"섹션 {i} 블록 {j} (table): rows[{r}]는 배열이어야 합니다")
            if len(row) > len(b["headers"]):
                raise ValueError(f"섹션 {i} 블록 {j} (table): rows[{r}] 셀 수가 headers 수({len(b['headers'])})를 초과합니다")
    elif k == "data":
        for si, sr in enumerate(b.get("series", []), 1):
            if not (isinstance(sr, dict) and isinstance(sr.get("name"), str)
                    and isinstance(sr.get("values"), list)):
                raise ValueError(f"섹션 {i} 블록 {j} (data): series[{si}]는 name(문자열)·values(배열)을 가져야 합니다")
            for v in sr["values"]:
                if isinstance(v, bool) or not isinstance(v, (int, float)):
                    raise ValueError(f"섹션 {i} 블록 {j} (data): series[{si}].values에 숫자가 아닌 값이 있습니다: {v!r}")


REQUIRED_SECTIONS = {
    "header": {"title"},
    "overview": {"title", "items"},
    "section": {"title", "blocks"},
    "conclusion": {"title"},
}
REQUIRED_BLOCKS = {
    "prose": {"heading", "paragraphs"},
    "table": {"heading", "headers", "rows"},
    "data": {"heading", "categories", "series"},
}


# ---------------------------------------------------------------- 파일명/채번

def _sanitize_title(title: str) -> str:
    """제목 → 파일명 접두어. Windows 금지 문자·glob 메타문자 [ ] 제거, 공백→_. (build_ppt.py와 동일 규칙)"""
    for ch in '\\/:*?"<>|[]\n\r\t':
        title = title.replace(ch, " ")
    return "_".join(title.split()).strip(" ._") or "제목없음"


def _next_version(out_dir: Path, prefix: str, ext: str) -> int:
    """out_dir에서 같은 접두어·확장자의 최대 vNN을 찾아 +1 (확장자별 독립 시퀀스)."""
    mx = 0
    for f in out_dir.glob(f"{prefix}_v*.{ext}"):
        m = re.search(rf"_v(\d+)\.{re.escape(ext)}$", f.name, re.IGNORECASE)
        if m:
            mx = max(mx, int(m.group(1)))
    return mx + 1


def _output_filename(meta: dict, out_dir: Path, ext: str) -> str:
    """output_name이 있으면 그대로(하위 호환, 확장자는 스크립트가 붙임), 없으면 <제목>_vNN.<ext>."""
    if meta.get("output_name"):
        name = str(meta["output_name"])
        if name.endswith("." + ext):
            name = name[: -len(ext) - 1]
        return f"{name}.{ext}"
    prefix = _sanitize_title(meta.get("title", ""))
    return f"{prefix}_v{_next_version(out_dir, prefix, ext):02d}.{ext}"


def _esc(s) -> str:
    return _html.escape(str(s), quote=False)


def _doc_title(doc) -> str:
    """문서 제목: header 섹션의 title 우선, 없으면 meta.title (md/html/docx 공통)."""
    header = next((s for s in doc.get("sections", []) if s.get("type") == "header"), {})
    return header.get("title") or doc.get("meta", {}).get("title", "")


# ---------------------------------------------------------------- markdown

def render_md(doc) -> str:
    meta = doc.get("meta", {})
    lines = [f"# {_doc_title(doc)}"]
    meta_line = next((s.get("meta_line", "") for s in doc.get("sections", [])
                      if s["type"] == "header" and s.get("meta_line")), "")
    if not meta_line:
        meta_line = header_line_join(meta)
    if meta_line:
        lines.append("")
        lines.append(f"> {meta_line}")

    for s in doc.get("sections", []):
        t = s["type"]
        if t == "header":
            continue  # H1에서 이미 표현
        if t == "overview":
            lines.append("")
            lines.append(f"## {s.get('title', '개요')}")
            for it in s.get("items", []):
                label = it.get("label", "")
                body = it.get("body", "")
                prefix = f"{it.get('no', '')} " if it.get("no") else ""
                label_part = f"**{label}** — " if label else ""
                lines.append(f"- {prefix}{label_part}{body}")
        elif t == "section":
            no = s.get("no", "")
            heading = f"## {no}. {s.get('title', '')}" if no else f"## {s.get('title', '')}"
            lines.append("")
            lines.append(heading)
            lead = s.get("lead", "")
            if lead:
                lines.append("")
                lines.append(lead)
            for b in s.get("blocks", []):
                lines.append("")
                lines.append(f"### {b.get('heading', '')}")
                if b["kind"] == "prose":
                    for p in b.get("paragraphs", []):
                        lines.append("")
                        lines.append(p)
                elif b["kind"] == "table":
                    lines.append("")
                    lines.append(_md_table(b.get("headers", []), b.get("rows", [])))
                elif b["kind"] == "data":
                    lines.append("")
                    lines.append(_md_data_table(b))
                    prose = b.get("prose", "")
                    if prose:
                        lines.append("")
                        lines.append(prose)
            src = s.get("source", "")
            if src:
                lines.append("")
                lines.append(f"> {src}")
        elif t == "conclusion":
            lines.append("")
            lines.append(f"## {s.get('title', '결론')}")
            for p in s.get("paragraphs", []):
                lines.append("")
                lines.append(p)
            reqs = s.get("requests", [])
            if reqs:
                lines.append("")
                for r in reqs:
                    lines.append(f"- 요청: {r}")
            note = s.get("note", "")
            if note:
                lines.append("")
                lines.append(f"> {note}")
    lines.append("")
    return "\n".join(lines)


def _md_cell(v) -> str:
    """셀 내 `|`·개행을 마크다운 표에서 안전하게 치환."""
    return str(v).replace("|", "\\|").replace("\r\n", "<br>").replace("\n", "<br>")


def _md_table(headers, rows):
    out = ["| " + " | ".join(_md_cell(h) for h in headers) + " |",
           "| " + " | ".join("---" for _ in headers) + " |"]
    out += ["| " + " | ".join(_md_cell(c) for c in r) + " |" for r in rows]
    return "\n".join(out)


def _md_data_table(block):
    cats = block.get("categories", [])
    series = block.get("series", [])
    unit = block.get("unit", "")
    head = [unit or "항목"] + [sr["name"] for sr in series]
    rows = []
    for i, cat in enumerate(cats):
        rows.append([cat] + [str(sr.get("values", [])[i]) if i < len(sr.get("values", [])) else "" for sr in series])
    return _md_table(head, rows)


def header_line_split(meta):
    return [meta.get("subtitle", ""), meta.get("doc_type", ""),
            meta.get("department", ""), meta.get("company", ""), meta.get("date", "")]


def header_line_join(meta):
    return " · ".join(x for x in header_line_split(meta) if x)


# ---------------------------------------------------------------- html

_HTML_CSS = """
:root {{
  --navy: #{navy}; --blue: #{blue}; --light-blue: #{light_blue};
  --bg: #{bg}; --bg-soft: #{bg_soft}; --text: #{text}; --text-sub: #{text_sub};
  --line: #{line}; --accent: #{accent}; --meta-line: #{meta_line};
}}
* {{ box-sizing: border-box; margin: 0; padding: 0; }}
body {{
  font-family: '{font}', '맑은 고딕', sans-serif;
  color: var(--text); background: var(--bg);
  max-width: 860px; margin: 0 auto; padding: 0 24px 64px;
  line-height: 1.75; font-size: 16px;
}}
.doc-header {{ background: var(--navy); color: #fff; margin: 0 -24px 40px; padding: 56px 48px 44px; }}
.doc-header h1 {{ font-size: 34px; line-height: 1.3; margin-bottom: 10px; }}
.doc-header .subtitle {{ color: var(--light-blue); font-size: 17px; }}
.doc-header .meta-line {{ color: var(--meta-line); font-size: 13px; margin-top: 18px; }}
h2 {{ font-size: 23px; color: var(--navy); margin: 48px 0 4px; padding-bottom: 6px;
      border-bottom: 2px solid var(--blue); }}
.overview .item {{ display: flex; gap: 14px; background: var(--bg-soft);
  border-radius: 6px; padding: 12px 18px; margin-top: 10px; }}
.overview .no {{ color: var(--blue); font-weight: 700; min-width: 26px; }}
.overview .label {{ font-weight: 700; min-width: 90px; }}
.lead {{ color: var(--text); margin-top: 14px; }}
h3 {{ font-size: 17px; color: var(--blue); margin: 26px 0 8px; }}
p {{ margin-top: 8px; }}
table {{ width: 100%; border-collapse: collapse; margin-top: 10px; font-size: 14px; }}
th {{ background: var(--navy); color: #fff; font-weight: 700; }}
th, td {{ border: 1px solid var(--line); padding: 8px 12px; text-align: left; }}
tbody tr:nth-child(even) {{ background: var(--bg-soft); }}
.bars {{ margin-top: 12px; }}
.bar-row {{ display: grid; grid-template-columns: 120px 1fr 84px; gap: 10px;
  align-items: center; padding: 4px 0; }}
.bar-label {{ font-size: 14px; color: var(--text-sub); text-align: right; }}
.bar-track {{ background: var(--bg-soft); border-radius: 3px; height: 18px; }}
.bar {{ height: 18px; border-radius: 3px; min-width: 2px; }}
.bar-value {{ font-size: 13px; color: var(--text); }}
.series-name {{ font-weight: 700; color: var(--blue); margin-top: 12px; }}
.source {{ color: var(--text-sub); font-size: 13px; margin-top: 14px; }}
.conclusion {{ background: var(--bg-soft); border-left: 4px solid var(--navy);
  padding: 20px 24px; margin-top: 16px; border-radius: 0 6px 6px 0; }}
.conclusion ul {{ margin: 10px 0 0 22px; }}
.conclusion .note {{ color: var(--text-sub); font-size: 14px; margin-top: 10px; }}
@media print {{ body {{ max-width: none; }} .doc-header {{ margin: 0 0 40px; }} }}
"""


def render_html(doc) -> str:
    meta = doc.get("meta", {})
    body = []
    for s in doc.get("sections", []):
        t = s["type"]
        if t == "header":
            meta_line = s.get("meta_line") or header_line_join(meta)  # md와 동일 폴백
            body.append(
                "<header class='doc-header'>"
                f"<h1>{_esc(s.get('title', ''))}</h1>"
                + (f"<div class='subtitle'>{_esc(s['subtitle'])}</div>" if s.get("subtitle") else "")
                + (f"<div class='meta-line'>{_esc(meta_line)}</div>" if meta_line else "")
                + "</header>")
        elif t == "overview":
            body.append(f"<h2>{_esc(s.get('title', '개요'))}</h2><div class='overview'>")
            for it in s.get("items", []):
                no = f"<span class='no'>{_esc(it['no'])}</span>" if it.get("no") else ""
                label = f"<span class='label'>{_esc(it['label'])}</span>" if it.get("label") else ""
                body.append(f"<div class='item'>{no}{label}<span>{_esc(it.get('body', ''))}</span></div>")
            body.append("</div>")
        elif t == "section":
            no = f"{s['no']}. " if s.get("no") else ""
            body.append(f"<h2>{_esc(no + s.get('title', ''))}</h2>")
            if s.get("lead"):
                body.append(f"<p class='lead'>{_esc(s['lead'])}</p>")
            for b in s.get("blocks", []):
                body.append(f"<h3>{_esc(b.get('heading', ''))}</h3>")
                if b["kind"] == "prose":
                    for p in b.get("paragraphs", []):
                        body.append(f"<p>{_esc(p)}</p>")
                elif b["kind"] == "table":
                    body.append(_html_table(b.get("headers", []), b.get("rows", [])))
                elif b["kind"] == "data":
                    body.extend(_html_data_bars(b))
            if s.get("source"):
                body.append(f"<p class='source'>{_esc(s['source'])}</p>")
        elif t == "conclusion":
            body.append(f"<h2>{_esc(s.get('title', '결론'))}</h2><div class='conclusion'>")
            for p in s.get("paragraphs", []):
                body.append(f"<p>{_esc(p)}</p>")
            reqs = s.get("requests", [])
            if reqs:
                body.append("<ul>" + "".join(f"<li>{_esc(r)}</li>" for r in reqs) + "</ul>")
            if s.get("note"):
                body.append(f"<div class='note'>{_esc(s['note'])}</div>")
            body.append("</div>")

    css = _HTML_CSS.format(navy=T.NAVY, blue=T.BLUE, light_blue=T.LIGHT_BLUE,
                           bg=T.BG, bg_soft=T.BG_SOFT, text=T.TEXT,
                           text_sub=T.TEXT_SUB, line=T.LINE, accent=T.ACCENT,
                           meta_line=META_LINE_COLOR, font=T.FONT)
    return ("<!doctype html>\n<html lang=\"ko\">\n<head>\n<meta charset=\"utf-8\">\n"
            "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">\n"
            f"<title>{_esc(_doc_title(doc))}</title>\n"
            f"<style>\n{css}</style>\n</head>\n<body>\n"
            + "\n".join(body) + "\n</body>\n</html>\n")


def _html_table(headers, rows):
    th = "".join(f"<th>{_esc(h)}</th>" for h in headers)
    trs = "".join("<tr>" + "".join(f"<td>{_esc(c)}</td>" for c in r) + "</tr>" for r in rows)
    return f"<table><thead><tr>{th}</tr></thead><tbody>{trs}</tbody></table>"


def _html_data_bars(block):
    """차트 데이터를 CSS 가로 막대로 렌더 (외부 리소스 0건)."""
    cats = block.get("categories", [])
    series = block.get("series", [])
    unit = block.get("unit", "")
    all_vals = [v for sr in series for v in sr.get("values", [])]
    vmax = max((float(v) for v in all_vals if v is not None), default=0)
    out = []
    if unit:
        out.append(f"<p style='color:{'#' + T.TEXT_SUB};font-size:13px'>단위: {_esc(unit)}</p>")
    out.append("<div class='bars'>")
    for si, sr in enumerate(series):
        color = "#" + SERIES_COLORS[si % len(SERIES_COLORS)]
        if len(series) > 1:
            out.append(f"<div class='series-name'>{_esc(sr.get('name', ''))}</div>")
        for ci, cat in enumerate(cats):
            vals = sr.get("values", [])
            val = vals[ci] if ci < len(vals) else 0
            pct = (float(val) / vmax * 100) if vmax > 0 and float(val) > 0 else 0
            out.append(
                "<div class='bar-row'>"
                f"<span class='bar-label'>{_esc(cat)}</span>"
                f"<div class='bar-track'><div class='bar' style='width:{pct:.1f}%;background:{color}'></div></div>"
                f"<span class='bar-value'>{_esc(val)}{_esc(unit)}</span></div>")
    out.append("</div>")
    if block.get("prose"):
        out.append(f"<p>{_esc(block['prose'])}</p>")
    return out


# ---------------------------------------------------------------- docx

def render_docx(doc, out_path: Path):
    try:
        from docx import Document
    except ImportError:
        raise SystemExit("python-docx가 필요합니다: pip install python-docx")
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    def set_style_font(style, size, bold=None, color=None):
        style.font.name = T.FONT
        style.font.size = Pt(size)
        if bold is not None:
            style.font.bold = bold
        if color:
            style.font.color.rgb = RGBColor.from_string(color)
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), T.FONT)  # python-docx는 font.name만으로 한글 글리프에 미적용

    def shade(cell, fill):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), fill)
        cell._tc.get_or_add_tcPr().append(shd)

    docx_doc = Document()
    set_style_font(docx_doc.styles["Normal"], 10.5)
    set_style_font(docx_doc.styles["Heading 1"], 20, bold=True, color=T.NAVY)
    set_style_font(docx_doc.styles["Heading 2"], 14, bold=True, color=T.NAVY)
    set_style_font(docx_doc.styles["Heading 3"], 11, bold=True, color=T.BLUE)

    meta = doc.get("meta", {})
    for s in doc.get("sections", []):
        t = s["type"]
        if t == "header":
            h = docx_doc.add_heading(s.get("title") or meta.get("title", ""), level=0)
            for r in h.runs:
                r.font.color.rgb = RGBColor.from_string(T.NAVY)
                r.font.name = T.FONT
            if s.get("subtitle"):
                p = docx_doc.add_paragraph(s["subtitle"])
                for r in p.runs:
                    r.font.color.rgb = RGBColor.from_string(T.TEXT_SUB)
            meta_line = s.get("meta_line") or header_line_join(meta)  # md/html과 동일 폴백
            if meta_line:
                p = docx_doc.add_paragraph(meta_line)
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor.from_string(T.TEXT_SUB)
        elif t == "overview":
            docx_doc.add_heading(s.get("title", "개요"), level=1)
            for it in s.get("items", []):
                no = f"{it['no']} " if it.get("no") else ""
                label = f"{it['label']} — " if it.get("label") else ""
                docx_doc.add_paragraph(f"{no}{label}{it.get('body', '')}")
        elif t == "section":
            no = f"{s['no']}. " if s.get("no") else ""
            docx_doc.add_heading(no + s.get("title", ""), level=1)
            if s.get("lead"):
                docx_doc.add_paragraph(s["lead"])
            for b in s.get("blocks", []):
                docx_doc.add_heading(b.get("heading", ""), level=2)
                if b["kind"] == "prose":
                    for p in b.get("paragraphs", []):
                        docx_doc.add_paragraph(p)
                elif b["kind"] == "table":
                    _docx_table(docx_doc, b.get("headers", []), b.get("rows", []), shade)
                elif b["kind"] == "data":
                    _docx_data_table(docx_doc, b, shade)
            if s.get("source"):
                p = docx_doc.add_paragraph(s["source"])
                for r in p.runs:
                    r.font.size = Pt(8.5)
                    r.font.color.rgb = RGBColor.from_string(T.TEXT_SUB)
        elif t == "conclusion":
            docx_doc.add_heading(s.get("title", "결론"), level=1)
            for p in s.get("paragraphs", []):
                docx_doc.add_paragraph(p)
            for r in s.get("requests", []):
                docx_doc.add_paragraph(f"요청: {r}", style="List Bullet")  # List Bullet이 마커를 붙임 (md는 "- 요청: ...")
            if s.get("note"):
                docx_doc.add_paragraph(s["note"])

    docx_doc.save(str(out_path))


def _docx_table(docx_doc, headers, rows, shade):
    from docx.shared import RGBColor

    table = docx_doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = str(h)
        shade(cell, T.NAVY)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor.from_string(T.BG)
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)


def _docx_data_table(docx_doc, block, shade):
    cats = block.get("categories", [])
    series = block.get("series", [])
    unit = block.get("unit", "")
    head = [unit or "항목"] + [sr["name"] for sr in series]
    rows = []
    for i, cat in enumerate(cats):
        rows.append([cat] + [str(sr.get("values", [])[i]) if i < len(sr.get("values", [])) else "" for sr in series])
    _docx_table(docx_doc, head, rows, shade)
    if block.get("prose"):
        docx_doc.add_paragraph(block["prose"])


# ---------------------------------------------------------------- 조립

def build(json_path: str, fmt: str, out_dir: str = ""):
    fmt = fmt.lower()
    if fmt not in FORMATS:
        raise ValueError(f"알 수 없는 포맷 '{fmt}' ({', '.join(FORMATS)} 중 하나여야 함)")
    jpath = Path(json_path)
    doc = json.loads(jpath.read_text(encoding="utf-8-sig"))
    sections = validate(doc)
    meta = doc.get("meta", {})

    base = Path(out_dir) if out_dir else jpath.resolve().parent.parent / "output"
    base.mkdir(parents=True, exist_ok=True)
    out = base / _output_filename(meta, base, fmt)

    if fmt == "md":
        out.write_text(render_md(doc), encoding="utf-8")
    elif fmt == "html":
        out.write_text(render_html(doc), encoding="utf-8")
    else:
        render_docx(doc, out)
    print(f"OK: {out} ({len(sections)} sections)")


if __name__ == "__main__":
    if len(sys.argv) not in (3, 4):
        print("사용법: python scripts/build_doc.py <report.json> <md|html|docx> [output_dir]")
        sys.exit(1)
    args = sys.argv[2:]
    fmts = [a for a in args if a.lower() in FORMATS]
    others = [a for a in args if a.lower() not in FORMATS]
    if len(others) > 1:
        print(f"인자를 해석할 수 없습니다: {', '.join(others)} "
              f"(포맷 {', '.join(FORMATS)} 중 1개 + 선택적 output_dir)")
        sys.exit(1)
    if not fmts:
        print(f"포맷 인자 필요: {', '.join(FORMATS)}")
        sys.exit(1)
    try:
        build(sys.argv[1], fmts[0], others[0] if others else "")
    except (ValueError, OSError) as e:  # JSONDecodeError·UnicodeDecodeError·FileNotFoundError 포함
        print(f"오류: {e}")
        sys.exit(1)